"""Genera la presentacion "Estado Ampliaciones" a partir de un corte.

Parte de la ultima presentacion real (assets/informe_direccion_plantilla.pptx)
y solo sustituye los datos: valores cacheados de los graficos, textos de
entradas/salidas y la fecha de portada. Asi se conservan intactos formato,
colores y maquetacion. Los graficos de la plantilla estaban vinculados al
Excel generador; ese vinculo se elimina y los datos quedan dentro del PPT.

Diapositivas que se rellenan (indice 1-based de la plantilla):
  1  portada (fecha)
  4  proyectos en curso + horas (ultimos 12 cortes DIRECCION)
  5  entradas / salidas respecto al corte DIRECCION anterior
  6  tartas por tamano de proyecto
  7  barras agrupacion de horas
  8  facturacion vs avance
  9  facturacion vs horas invertidas
  10 proyectos entregados (acumulado ano anterior vs actual)
  11 horas entregadas (acumulado ano anterior vs actual)
  12-16 objetivos kick-off / diseno / validacion / go-live / end
  17-18 anexo de hitos retrasados
El resto (otros temas, cierre) queda como en la plantilla.
"""

from __future__ import annotations

import copy
import io
import re
from datetime import date
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.util import Emu

from reports.informe_objetivos import (
    ESTADO_NARANJA,
    ESTADO_OK,
    ESTADO_ROJO,
    HITOS,
    mes_titulo,
)

TEMPLATE_PATH = Path(__file__).resolve().parent / "assets" / "informe_direccion_plantilla.pptx"

C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"c": C_NS}

HISTORY_POINTS = 12
SIZE_KEYS = ("s", "m", "l")
BILLING_KEYS = ("total", "s", "m", "l")
BILLING_SERIES = ("base", "facturado", "rojo", "verde")
MONTH_LABELS = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]


# ------------------------------------
# Graficos: sustitucion de valores cacheados
# ------------------------------------

def _excel_serial(value: date) -> int:
    return (value - date(1899, 12, 30)).days


def _c(tag: str) -> str:
    return f"{{{C_NS}}}{tag}"


def _fill_cache(container: etree._Element, values: list) -> None:
    """Rellena numRef/strRef (su cache) o numLit/strLit con `values`.
    Los None se omiten (punto vacio), igual que una celda en blanco."""
    data = None
    for child in container:
        tag = etree.QName(child).localname
        if tag in ("numRef", "strRef"):
            cache_tag = "numCache" if tag == "numRef" else "strCache"
            data = child.find(f"c:{cache_tag}", NS)
            if data is None:
                data = etree.SubElement(child, _c(cache_tag))
            break
        if tag in ("numLit", "strLit"):
            data = child
            break
    if data is None:
        raise ValueError("Serie de grafico sin datos cacheados")

    for pt in data.findall("c:pt", NS):
        data.remove(pt)
    count = data.find("c:ptCount", NS)
    if count is None:
        count = etree.Element(_c("ptCount"))
        format_code = data.find("c:formatCode", NS)
        if format_code is not None:
            format_code.addnext(count)
        else:
            data.insert(0, count)
    count.set("val", str(len(values)))

    anchor = count
    for index, value in enumerate(values):
        if value is None:
            continue
        pt = etree.Element(_c("pt"))
        pt.set("idx", str(index))
        etree.SubElement(pt, _c("v")).text = str(value)
        anchor.addnext(pt)
        anchor = pt


def _set_series_name(ser: etree._Element, name: str) -> None:
    tx = ser.find("c:tx", NS)
    if tx is None:
        return
    v = tx.find("c:v", NS)
    if v is not None:
        v.text = name
    else:
        _fill_cache(tx, [name])


def _unlink_chart(chart) -> None:
    """Quita el vinculo al Excel externo para que PowerPoint no pida actualizarlo."""
    chart_space = chart._chartSpace
    external = chart_space.find("c:externalData", NS)
    if external is None:
        return
    rel_id = external.get(f"{{{R_NS}}}id")
    chart_space.remove(external)
    if rel_id and rel_id in chart.part.rels:
        chart.part.drop_rel(rel_id)


def _update_chart(chart, categories: list, series: list[list], names: list[str] | None = None) -> None:
    sers = chart._chartSpace.findall(".//c:ser", NS)
    if len(sers) != len(series):
        raise ValueError(f"El grafico tiene {len(sers)} series y se esperaban {len(series)}")
    for index, (ser, values) in enumerate(zip(sers, series)):
        cat = ser.find("c:cat", NS)
        if cat is not None:
            _fill_cache(cat, categories)
        _fill_cache(ser.find("c:val", NS), values)
        if names:
            _set_series_name(ser, names[index])
    _unlink_chart(chart)


def _chart_by_name(slide, name: str):
    """La plantilla tiene copias vacias (sin series) con el mismo nombre; se usa la que tiene datos."""
    for shape in slide.shapes:
        if shape.name == name and getattr(shape, "has_chart", False) and shape.has_chart:
            if shape.chart._chartSpace.find(".//c:ser", NS) is not None:
                return shape.chart
    raise ValueError(f"No se encuentra el grafico '{name}' en la diapositiva")


# ------------------------------------
# Textos
# ------------------------------------

LEGAL_SUFFIX_RE = re.compile(
    r"[\s,.-]+(S\.?\s?A\.?\s?U\.?|S\.?\s?L\.?\s?U\.?|S\.?\s?A\.?|S\.?\s?L\.?|B\.?\s?V\.?|S\.?\s?P\.?\s?A\.?|SAS|LTDA?\.?|GMBH)\s*$",
    re.IGNORECASE,
)


def _short_name(item: dict) -> str:
    name = item.get("client") or item.get("project_name") or item.get("code") or ""
    previous = None
    while previous != name:
        previous = name
        name = LEGAL_SUFFIX_RE.sub("", name).strip(" ,.-")
    return name


def _fmt_hours(value: float) -> str:
    if abs(value - round(value)) < 0.05:
        return f"{round(value):.0f}"
    return f"{value:.1f}".replace(".", ",")


def _set_runs(paragraph, texts: list[str]) -> None:
    runs = paragraph.runs
    for run, text in zip(runs, texts):
        run.text = text
    for run in runs[len(texts):]:
        run.text = ""


def _clone_paragraph(prototype, texts: list[str]):
    element = copy.deepcopy(prototype._p)
    from pptx.text.text import _Paragraph

    paragraph = _Paragraph(element, prototype._parent)
    _set_runs(paragraph, texts)
    return element


def _rebuild_text(text_frame, blocks: list[tuple[object, list[str]]]) -> None:
    """Sustituye los parrafos por clones de los prototipos indicados."""
    elements = [_clone_paragraph(proto, texts) for proto, texts in blocks]
    tx_body = text_frame._txBody
    for p in list(tx_body.p_lst):
        tx_body.remove(p)
    for element in elements:
        tx_body.append(element)


def _shape_by_text(slide, startswith: str):
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip().upper().startswith(startswith):
            return shape
    raise ValueError(f"No se encuentra el cuadro de texto '{startswith}'")


def _tables(slide) -> list:
    return sorted((s for s in slide.shapes if s.has_table), key=lambda s: s.left)


def _set_cell(cell, text: str) -> None:
    paragraph = cell.text_frame.paragraphs[0]
    if paragraph.runs:
        _set_runs(paragraph, [text])
    else:
        paragraph.text = text


def _fill_entradas_salidas(slide, delta: dict | None) -> None:
    entradas_box = _shape_by_text(slide, "ENTRADAS")
    salidas_box = _shape_by_text(slide, "SALIDAS")
    tabla_entradas, tabla_salidas = _tables(slide)

    if delta is None:
        return

    # Prototipos de formato tomados de la plantilla
    ep = entradas_box.text_frame.paragraphs
    e_header, e_cambios, e_cambio_item = ep[0], ep[1], ep[2]
    e_section = next(p for p in ep if p.text.strip().startswith("Proyectos >400"))
    e_none = next(p for p in ep if p.text.strip() == "Ninguno")
    e_item = next(p for p in ep if p.level == 1 and p.runs and p.runs[0].font.size and p.runs[0].font.size.pt == 11)
    e_small = next(p for p in ep if p.text.strip().startswith("Proyectos <100"))

    sp = salidas_box.text_frame.paragraphs
    s_header = sp[0]
    s_section = next(p for p in sp if p.text.strip().startswith("Proyectos >400"))
    s_none = next(p for p in sp if p.text.strip() == "Ninguno")
    s_item = next(p for p in sp if p.level == 1 and p.text.strip() != "Ninguno")
    s_small = next(p for p in sp if p.text.strip().startswith("Proyectos <100"))

    entradas = delta["entradas"]
    salidas = delta["salidas"]

    def _items(tramo: list[dict], item_proto, none_proto) -> list:
        if not tramo:
            return [(none_proto, ["Ninguno"])]
        return [(item_proto, [f"{_short_name(i)} ({_fmt_hours(i['hours'])} h)"]) for i in tramo]

    blocks = [
        (e_header, ["ENTRADAS", ": ( ", f"{_fmt_hours(entradas['total_horas'])} ", "h)"]),
        (e_cambios, ["Cambio horas proyectos existentes: ", f"{_fmt_hours(entradas['cambios_horas'])} h "]),
    ]
    blocks += [(e_cambio_item, [f"{_short_name(i)} ({_fmt_hours(i['hours'])} h)"]) for i in entradas["cambios"]]
    blocks.append((e_section, ["Proyectos >400 h:"]))
    blocks += _items(entradas["tramos"]["l"], e_item, e_none)
    blocks.append((e_section, ["Proyectos entre 100 y 400 h: "]))
    blocks += _items(entradas["tramos"]["m"], e_item, e_none)
    blocks.append((e_small, ["Proyectos <100 horas"]))
    _rebuild_text(entradas_box.text_frame, blocks)

    blocks = [(s_header, ["SALIDAS", ": ", f"({_fmt_hours(salidas['total_horas'])} h)"])]
    blocks.append((s_section, ["Proyectos >400 h"]))
    blocks += _items(salidas["tramos"]["l"], s_item, s_none)
    blocks.append((s_section, ["Proyectos entre 100 y 400 h:"]))
    blocks += _items(salidas["tramos"]["m"], s_item, s_none)
    blocks.append((s_small, ["Proyectos <100 horas"]))
    _rebuild_text(salidas_box.text_frame, blocks)

    for table_shape, tramo in ((tabla_entradas, entradas["tramos"]["s"]), (tabla_salidas, salidas["tramos"]["s"])):
        table = table_shape.table
        _set_cell(table.cell(0, 1), str(len(tramo)))
        _set_cell(table.cell(1, 1), _fmt_hours(sum(i["hours"] for i in tramo)))


# ------------------------------------
# Objetivos (12-16) y hitos retrasados (17-18)
# ------------------------------------

OBJETIVOS_SLIDES = {"K": 11, "D": 12, "V": 13, "G": 14, "E": 15}
RETRASADOS_SLIDES = (16, 17)
RETRASADOS_TITULOS = {"K": "KICK-OFF", "D": "DISEÑO", "V": "VALIDACIÓN", "G": "GO-LIVE", "E": "END"}

# Colores de la plantilla para las listas de proyectos
COLOR_ESTADO = {
    ESTADO_OK: RGBColor(0x00, 0x58, 0x92),
    ESTADO_NARANJA: RGBColor(0xFF, 0x82, 0x05),
    ESTADO_ROJO: RGBColor(0xFF, 0x00, 0x00),
}
TEXT = "text"      # color de texto del tema
GREY = "grey"      # gris de "Ninguno"

# Maquetacion del anexo (EMU), tomada de la plantilla
RET_TOP = (1312223, 1429085)
RET_BOTTOM = 6300000
RET_HEADER_TO_BODY = 333760
RET_LINE = 168000
RET_CHARS_PER_LINE = 150
RET_GAP = 120000


def _prototype_paragraph(table):
    """Primer parrafo con texto de la tabla en una celda de lista (fila 2), como molde."""
    for row_index in (2, 3, 1):
        for cell in list(table.rows[row_index].cells)[1:]:
            for paragraph in cell.text_frame.paragraphs:
                if paragraph.runs and paragraph.runs[0].text.strip():
                    return paragraph
    raise ValueError("Tabla de objetivos sin texto de ejemplo")


def _color_run(run, color) -> None:
    if color == TEXT:
        run.font.color.theme_color = MSO_THEME_COLOR.TEXT_1
    elif color == GREY:
        run.font.color.theme_color = MSO_THEME_COLOR.BACKGROUND_2
    else:
        run.font.color.rgb = color


def _write_cell(cell, prototype, lines: list[list[tuple[str, object]]]) -> None:
    """Escribe en la celda una linea por elemento; cada linea es [(texto, color)]."""
    from pptx.text.text import _Paragraph

    tx_body = cell.text_frame._txBody
    for p in list(tx_body.p_lst):
        tx_body.remove(p)
    for parts in lines:
        element = copy.deepcopy(prototype._p)
        tx_body.append(element)
        paragraph = _Paragraph(element, cell.text_frame)
        runs = paragraph.runs
        for extra in runs[1:]:
            element.remove(extra._r)
        base = runs[0]._r
        for index, (text, color) in enumerate(parts):
            r = base if index == 0 else copy.deepcopy(base)
            if index:
                element.append(r)
            run = paragraph.runs[index]
            run.text = text
            _color_run(run, color)
        # el fin de parrafo debe quedar al final
        end = element.find("{http://schemas.openxmlformats.org/drawingml/2006/main}endParaRPr")
        if end is not None:
            element.remove(end)
            element.append(end)


def _fill_objetivos(slide, tabla: dict) -> None:
    table = next(s for s in slide.shapes if s.has_table).table
    prototype = copy.deepcopy(_prototype_paragraph(table))
    meses = tabla["meses"]
    for column, ym in enumerate(meses, start=1):
        header = table.cell(0, column).text_frame.paragraphs[0]
        _set_runs(header, [mes_titulo(ym)])

        small = tabla["celdas"][(ym, "s")]
        if not small:
            lines = [[("Ninguno", GREY)]]
        elif ym <= tabla["mes_corte"]:
            ok = sum(1 for e in small if e["estado"] == ESTADO_OK)
            lines = [[(f"{len(small)} (", TEXT), (f"conseguidos {ok}", COLOR_ESTADO[ESTADO_OK]), (")", TEXT)]]
        else:
            lines = [[(str(len(small)), TEXT)]]
        _write_cell(table.cell(1, column), prototype, lines)

        for row, tramo in ((2, "m"), (3, "l")):
            entries = tabla["celdas"][(ym, tramo)]
            if entries:
                lines = [[(e["nombre"], COLOR_ESTADO.get(e["estado"], TEXT))] for e in entries]
            else:
                lines = [[("Ninguno", GREY)]]
            _write_cell(table.cell(row, column), prototype, lines)


def _retrasado_texto(item: dict) -> str:
    motivo = (item.get("motivo") or "").strip()
    texto = f"{item['nombre']} ({_fmt_hours(item['horas'])} h)"
    return f"{texto}: {motivo}" if motivo else texto


def _body_height(texts: list[str]) -> int:
    lines = sum(max(1, -(-len(t) // RET_CHARS_PER_LINE)) for t in texts)
    return lines * RET_LINE + RET_LINE // 2


def _fill_retrasados(prs, retrasados: dict[str, list[dict]]) -> None:
    slides = [prs.slides[i] for i in RETRASADOS_SLIDES]
    # Moldes: titulo de hito (16 pt azul) y cuerpo (11 pt) de la plantilla
    header_proto = body_proto = None
    for slide in slides:
        for shape in list(slide.shapes):
            if not shape.has_text_frame or not shape.name.startswith("CuadroTexto") or shape.name == "CuadroTexto 137":
                continue
            text = shape.text_frame.text.strip()
            is_header = text.upper() in RETRASADOS_TITULOS.values()
            if is_header and header_proto is None:
                header_proto = copy.deepcopy(shape._element)
            elif not is_header and body_proto is None and text:
                body_proto = copy.deepcopy(shape._element)
            shape._element.getparent().remove(shape._element)
    if header_proto is None or body_proto is None:
        raise ValueError("La plantilla no tiene los cuadros de texto del anexo de retrasados")

    groups = [(key, retrasados[key]) for key, *_ in HITOS if retrasados.get(key)]
    if not groups:
        groups = [(None, [])]

    # Bloques de como mucho lo que cabe en una diapositiva: un hito con muchos
    # proyectos se parte y continua en la siguiente con su mismo titulo.
    max_body = RET_BOTTOM - RET_TOP[1] - RET_HEADER_TO_BODY
    blocks = []
    for key, items in groups:
        texts = [_retrasado_texto(i) for i in items] or ["Sin hitos retrasados de 100 h o más."]
        chunk: list[str] = []
        for text in texts:
            if chunk and _body_height(chunk + [text]) > max_body:
                blocks.append((key, chunk))
                chunk = []
            chunk.append(text)
        blocks.append((key, chunk))

    blank = copy.deepcopy(slides[1]._element)  # anexo vacio, para duplicarlo si hace falta
    slide_index = 0
    y = RET_TOP[0]
    used = [False] * len(slides)
    for key, texts in blocks:
        height = RET_HEADER_TO_BODY + _body_height(texts)
        if y + height > RET_BOTTOM and used[slide_index]:
            slide_index += 1
            if slide_index == len(slides):
                slides.append(_duplicate_slide(prs, slides[-1], blank))
                used.append(False)
            y = RET_TOP[1]
        slide = slides[slide_index]
        used[slide_index] = True

        header = copy.deepcopy(header_proto)
        slide.shapes._spTree.append(header)
        header_shape = slide.shapes[-1]
        header_shape.top = Emu(y)
        _set_runs(header_shape.text_frame.paragraphs[0], [RETRASADOS_TITULOS[key] if key else "HITOS RETRASADOS"])

        body = copy.deepcopy(body_proto)
        slide.shapes._spTree.append(body)
        body_shape = slide.shapes[-1]
        body_shape.top = Emu(y + RET_HEADER_TO_BODY)
        body_shape.height = Emu(_body_height(texts))
        prototype = next(p for p in body_shape.text_frame.paragraphs if p.runs)
        _rebuild_text(body_shape.text_frame, [(prototype, [t]) for t in texts])
        y += height + RET_GAP

    for slide, was_used in zip(slides, used):
        if not was_used:
            _delete_slide(prs, slide)


def _duplicate_slide(prs, after, blank_element):
    """Nueva diapositiva de anexo (copia de `blank_element`) justo despues de `after`."""
    new = prs.slides.add_slide(after.slide_layout)
    sp_tree = new.shapes._spTree
    for shape in list(new.shapes):
        sp_tree.remove(shape._element)
    for child in blank_element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}spTree"):
        if etree.QName(child).localname in ("nvGrpSpPr", "grpSpPr"):
            continue
        sp_tree.append(copy.deepcopy(child))
    slide_ids = prs.slides._sldIdLst
    new_id = slide_ids[-1]
    position = next(i for i, sid in enumerate(slide_ids) if prs.slides.get(int(sid.get("id"))) is after)
    slide_ids.remove(new_id)
    slide_ids.insert(position + 1, new_id)
    return new


def _delete_slide(prs, slide) -> None:
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        if prs.slides.get(int(slide_id.get("id"))) is slide:
            prs.part.drop_rel(slide_id.rId)
            slide_ids.remove(slide_id)
            return


# ------------------------------------
# Presentacion completa
# ------------------------------------

def generar_presentacion(
    corte: dict,
    historico: list[dict],
    delta: dict | None,
    entregas_anterior: dict | None,
    objetivos: dict | None = None,
) -> bytes:
    """`historico`: cortes DIRECCION hasta el corte (incluido), ordenados por fecha.
    `delta`: resultado de comparar_cortes contra el corte DIRECCION anterior (o None).
    `entregas_anterior`: entregas acumuladas del ano anterior {"horas", "proyectos"}.
    `objetivos`: resultado de informe_objetivos.datos_objetivos (o None para dejar
    las diapositivas 12-18 como en la plantilla)."""
    prs = Presentation(str(TEMPLATE_PATH))
    slides = prs.slides
    metricas = corte["metricas"]
    fecha: date = corte["fecha_corte"]

    # 1. Portada
    portada = _shape_by_text(slides[0], "20")
    _set_runs(portada.text_frame.paragraphs[0], [fecha.isoformat()])

    # 4. Historico de proyectos y horas
    puntos = [c for c in historico if c["metricas"].get("proyectos") and c["metricas"].get("horas")][-HISTORY_POINTS:]
    fechas = [_excel_serial(c["fecha_corte"]) for c in puntos]
    _update_chart(
        _chart_by_name(slides[3], "Chart 1"), fechas,
        [[c["metricas"]["proyectos"]["totales"] for c in puntos],
         [c["metricas"]["proyectos"]["activos"] for c in puntos]],
    )
    _update_chart(
        _chart_by_name(slides[3], "cht_Graf_02_Horas"), fechas,
        [[c["metricas"]["horas"][key] for c in puntos] for key in ("totales", "pendientes", "ejecutables")],
    )

    # 5. Entradas / salidas
    _fill_entradas_salidas(slides[4], delta)

    # 6-7. Tramos por tamano de proyecto
    tramos = metricas["tramos"]
    _update_chart(
        _chart_by_name(slides[5], "cht_Graf_03_TrabajoHoras"),
        ["Horas proyectos <= 100", "Horas proyectos > 100 y < 400", "Horas proyectos >= 400"],
        [[tramos[k]["total"] for k in SIZE_KEYS]],
    )
    _update_chart(
        _chart_by_name(slides[5], "cht_Graf_04_TrabajoPendiente"),
        ["Horas pendientes proyectos <= 100", "Horas pendientes proyectos > 100 y < 400",
         "Horas pendientes proyectos >= 400"],
        [[tramos[k]["pendiente"] for k in SIZE_KEYS]],
    )
    horas = metricas["horas"]
    general = {"total": horas["totales"], "pendiente": horas["pendientes"], "ejecutable": horas["ejecutables"]}
    _update_chart(
        _chart_by_name(slides[6], "cht_Graf_05_BarrasAgrupacion"),
        ["<=100", "100-400", ">=400", "General"],
        [[tramos[k][serie] for k in SIZE_KEYS] + [general[serie]] for serie in ("total", "pendiente", "ejecutable")],
    )

    # 8-9. Facturacion
    for slide_index, chart_name, key in (
        (7, "cht_Graf_07_AvanceVsFacturacion", "avance_vs_facturacion"),
        (8, "cht_Graf_06_HorasVsFacturacion", "horas_vs_facturacion"),
    ):
        data = metricas.get(key)
        if data:
            _update_chart(
                _chart_by_name(slides[slide_index], chart_name),
                ["Total", "<=100", "100-400", ">=400"],
                [[data[group][serie] for group in BILLING_KEYS] for serie in BILLING_SERIES],
            )

    # 10-11. Entregas acumuladas
    entregas = metricas.get("entregas")
    if entregas:
        anterior = entregas_anterior or {"horas": [None] * 12, "proyectos": [None] * 12}
        names = [str(fecha.year - 1), str(fecha.year)]
        _update_chart(
            _chart_by_name(slides[9], "cht_Graf_09_EntregasProyectos"), MONTH_LABELS,
            [anterior["proyectos"], entregas["proyectos"]], names,
        )
        _update_chart(
            _chart_by_name(slides[10], "cht_Graf_08_EntregasHoras"), MONTH_LABELS,
            [anterior["horas"], entregas["horas"]], names,
        )

    # 12-18. Objetivos por hito y anexo de retrasados (las ultimas, porque
    # borrar una diapositiva sobrante cambia los indices)
    if objetivos and objetivos.get("objetivos"):
        for key, slide_index in OBJETIVOS_SLIDES.items():
            _fill_objetivos(slides[slide_index], objetivos["objetivos"][key])
        _fill_retrasados(prs, objetivos["retrasados"])

    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()
